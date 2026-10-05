import io
import os
import sqlite3
import time
from datetime import datetime
from typing import List, Literal

import streamlit as st
from google import genai
from google.genai import types
from pydantic import BaseModel, Field

try:
    from docx import Document
except Exception:
    Document = None

st.set_page_config(page_title="M-MIS Study AI", page_icon="🎓", layout="wide")

# ============================== CONFIG ==============================
MODEL_NAME = "gemini-3.8-flash"
MAX_FILES = 8
SUPPORTED_TYPES = ["pdf", "txt", "md", "csv", "png", "jpg", "jpeg", "webp", "docx"]
DB_PATH = "m_mis_study.db"

SYSTEM_INSTRUCTION = """
You are M-MIS Study AI, a university study assistant specialized in Management Information Systems (MIS), Business Administration, Management, Information Systems, and Business Technology.

ACADEMIC ACCURACY RULES — VERY IMPORTANT:
1. Uploaded course material is the primary source for material-specific questions.
2. Never invent a fact and claim it came from the uploaded material.
3. If a requested fact is not clearly supported by the uploaded material, say exactly: "This information is not clearly stated in the uploaded material."
4. Clearly distinguish source facts, reasonable inference, and general academic knowledge.
5. Mention a page/slide/section only when it is reliably identifiable. Never fabricate page numbers.
6. Preserve official definitions and their meaning. Simplify only after preserving the meaning.
7. Keep important English academic terms and explain them in Arabic when useful.
8. For summaries, preserve definitions, classifications, relationships, processes, formulas, and important examples.
9. Generated questions must be answerable from the selected uploaded material unless general-knowledge questions are explicitly requested.
10. If the source is ambiguous, incomplete, or contradictory, say so instead of guessing.
11. Never present uncertainty as certainty.
12. Be concise enough for studying, but complete enough to be useful for exams.
13. Default language is Egyptian/Modern Arabic as appropriate, while retaining important English terminology.
"""

# ============================== MODELS ==============================
class QuestionItem(BaseModel):
    question: str
    question_type: Literal["MCQ", "True/False", "Fill in the blank", "Essay"]
    difficulty: Literal["Easy", "Medium", "Hard"]
    options: List[str] = Field(default_factory=list)
    correct_answer: str
    explanation: str
    source: str
    topic: str = "General"


class QuestionSet(BaseModel):
    questions: List[QuestionItem]


class FlashcardItem(BaseModel):
    front: str
    back: str
    source: str
    topic: str = "General"


class FlashcardSet(BaseModel):
    cards: List[FlashcardItem]


class TopicSet(BaseModel):
    topics: List[str]


# ============================== DATABASE ==============================
def db_conn():
    return sqlite3.connect(DB_PATH, check_same_thread=False)


def init_db():
    con = db_conn()
    cur = con.cursor()
    cur.execute("""CREATE TABLE IF NOT EXISTS quiz_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        mode TEXT NOT NULL,
        score INTEGER NOT NULL,
        total INTEGER NOT NULL,
        percent INTEGER NOT NULL,
        topics TEXT,
        created_at TEXT NOT NULL
    )""")
    cur.execute("""CREATE TABLE IF NOT EXISTS study_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        action TEXT NOT NULL,
        material_count INTEGER NOT NULL,
        created_at TEXT NOT NULL
    )""")
    con.commit()
    con.close()


def save_result(mode, score, total, topics):
    con = db_conn()
    con.execute(
        "INSERT INTO quiz_history(mode,score,total,percent,topics,created_at) VALUES(?,?,?,?,?,?)",
        (mode, score, total, round(score / total * 100) if total else 0, ", ".join(topics[:10]), datetime.now().isoformat(timespec="seconds")),
    )
    con.commit()
    con.close()


def save_action(action, material_count):
    con = db_conn()
    con.execute(
        "INSERT INTO study_log(action,material_count,created_at) VALUES(?,?,?)",
        (action, material_count, datetime.now().isoformat(timespec="seconds")),
    )
    con.commit()
    con.close()


def get_history(limit=10):
    con = db_conn()
    rows = con.execute(
        "SELECT mode,score,total,percent,created_at FROM quiz_history ORDER BY id DESC LIMIT ?", (limit,)
    ).fetchall()
    con.close()
    return rows


init_db()

# ============================== UI ==============================
st.markdown(r"""
<style>
.stApp{background:radial-gradient(circle at 15% 0%,rgba(99,91,255,.13),transparent 28%),#080b14;color:#f5f7ff}
.block-container{max-width:1400px;padding-top:1.2rem;padding-bottom:4rem}
section[data-testid="stSidebar"]{background:#0d1220;border-right:1px solid rgba(255,255,255,.08)}
.brand{display:flex;gap:12px;align-items:center}.brand-icon{width:44px;height:44px;border-radius:14px;display:flex;align-items:center;justify-content:center;background:linear-gradient(135deg,#635bff,#8b5cf6);font-size:22px}.brand-title{font-size:20px;font-weight:800}.brand-subtitle{font-size:11px;color:#8d96aa}
.hero{padding:28px;border:1px solid rgba(255,255,255,.08);border-radius:24px;background:linear-gradient(135deg,rgba(99,91,255,.20),rgba(15,23,42,.55));box-shadow:0 18px 60px rgba(0,0,0,.22);margin-bottom:20px}.hero-kicker{color:#a5b4fc;font-weight:700;font-size:13px;text-transform:uppercase}.hero h1{margin:6px 0 0;font-size:clamp(30px,5vw,52px);line-height:1.05;letter-spacing:-1.8px}.hero p{color:#aeb7ca;max-width:900px;font-size:16px;line-height:1.7}
.card{border:1px solid rgba(255,255,255,.08);border-radius:20px;padding:20px;background:rgba(16,21,37,.72);box-shadow:0 12px 40px rgba(0,0,0,.15)}.feature-card{min-height:140px}.feature-icon{font-size:25px}.feature-title{font-size:17px;font-weight:800;margin:7px 0}.feature-text{color:#9da7bb;font-size:13px;line-height:1.55}.section-title{font-size:23px;font-weight:800;margin:26px 0 12px}.status{display:inline-flex;align-items:center;gap:7px;padding:7px 11px;border-radius:999px;background:rgba(34,197,94,.10);border:1px solid rgba(34,197,94,.18);color:#86efac;font-size:12px;font-weight:700}.status-dot{width:7px;height:7px;border-radius:50%;background:#4ade80}
.material-chip{display:flex;align-items:center;gap:10px;padding:10px 12px;border-radius:12px;background:rgba(255,255,255,.045);border:1px solid rgba(255,255,255,.06);margin-bottom:8px;font-size:12px}.material-chip b{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.quiz-question{padding:22px;border-radius:18px;border:1px solid rgba(255,255,255,.08);background:rgba(17,24,39,.75);margin-bottom:14px}.quiz-number{color:#a5b4fc;font-size:12px;font-weight:800;text-transform:uppercase}.quiz-text{font-size:20px;line-height:1.55;font-weight:750}
.score-card{text-align:center;padding:30px;border-radius:24px;background:linear-gradient(135deg,rgba(99,91,255,.18),rgba(15,23,42,.75));border:1px solid rgba(99,91,255,.25)}.score-number{font-size:58px;font-weight:900;line-height:1}.source-badge{display:inline-block;margin-top:8px;padding:5px 9px;border-radius:8px;background:rgba(59,130,246,.10);border:1px solid rgba(59,130,246,.18);color:#93c5fd;font-size:11px}
.metric{padding:18px;border-radius:18px;border:1px solid rgba(255,255,255,.07);background:rgba(255,255,255,.035)}.metric-number{font-size:28px;font-weight:900}.metric-label{color:#8d96aa;font-size:12px}
@media(max-width:768px){.block-container{padding:.7rem .65rem 3rem}.hero{padding:20px;border-radius:18px}.hero h1{font-size:32px}.section-title{font-size:20px}.quiz-text{font-size:18px}}
</style>
""", unsafe_allow_html=True)

# ============================== API ==============================
api_key = st.secrets.get("GEMINI_API_KEY") or os.getenv("GEMINI_API_KEY")
if not api_key:
    st.error("⚠️ GEMINI_API_KEY غير موجود. أضيفيه في Streamlit Secrets أو كـ environment variable.")
    st.stop()

try:
    client = genai.Client(api_key=api_key, http_options={"api_version": "v1"})
except Exception as exc:
    st.error(f"تعذر تهيئة Gemini: {exc}")
    st.stop()

# ============================== STATE ==============================
def init_state():
    defaults = {
        "materials": {}, "selected_materials": [], "chat_messages": [],
        "chat_previous_id": None,
        "quiz_questions": [], "quiz_answers": {}, "quiz_submitted": False,
        "quiz_score": None, "flashcards": [], "flashcard_index": 0,
        "flashcard_revealed": False, "exam_questions": [], "exam_answers": {},
        "exam_submitted": False, "exam_score": None, "exam_started_at": None,
        "study_actions": 0,
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v.copy() if isinstance(v, (dict, list)) else v


init_state()


def reset_study_state():
    for k, v in {
        "chat_messages": [], "chat_previous_id": None,
        "quiz_questions": [], "quiz_answers": {}, "quiz_submitted": False,
        "quiz_score": None, "flashcards": [], "flashcard_index": 0,
        "flashcard_revealed": False, "exam_questions": [], "exam_answers": {},
        "exam_submitted": False, "exam_score": None, "exam_started_at": None,
    }.items():
        st.session_state[k] = v.copy() if isinstance(v, (dict, list)) else v


def selected_items():
    return [st.session_state.materials[n] for n in st.session_state.selected_materials if n in st.session_state.materials]


def interaction_input(prompt):
    items = []
    for item in selected_items():
        f = item["gemini_file"]
        items.append({"type": "document", "uri": f.uri, "mime_type": f.mime_type})
    items.append({"type": "text", "text": prompt})
    return items


def ask(prompt, schema=None, temperature=0.1, continue_chat=False):
    kwargs = {
        "model": MODEL_NAME,
        "input": interaction_input(prompt),
        "system_instruction": SYSTEM_INSTRUCTION,
        "generation_config": {"temperature": temperature},
    }
    if schema:
        kwargs["response_format"] = {
            "type": "text",
            "mime_type": "application/json",
            "schema": schema.model_json_schema(),
        }
    if continue_chat and st.session_state.chat_previous_id:
        kwargs["previous_interaction_id"] = st.session_state.chat_previous_id
        # Previous interaction already contains the previous documents; still include the
        # selected docs in the current input so changing context remains explicit.
    interaction = client.interactions.create(**kwargs)
    if getattr(interaction, "status", "completed") not in ("completed", "in_progress"):
        raise RuntimeError(f"Gemini interaction status: {interaction.status}")
    text = interaction.output_text or ""
    if not text:
        raise RuntimeError("Gemini returned an empty response.")
    if continue_chat:
        st.session_state.chat_previous_id = interaction.id
    return schema.model_validate_json(text) if schema else text


def upload_bytes(name, data, mime):
    return client.files.upload(
        file=io.BytesIO(data),
        config=types.UploadFileConfig(display_name=name, mime_type=mime),
    )


def process_upload(up):
    data = up.getvalue()
    name = up.name
    mime = up.type or "application/octet-stream"
    if name.lower().endswith(".docx"):
        if Document is None:
            raise RuntimeError("python-docx غير مثبت. ثبتي requirements الجديدة.")
        doc = Document(io.BytesIO(data))
        text_parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                text_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        data = ("\n".join(text_parts)).encode("utf-8")
        name = name + " (extracted text)"
        mime = "text/plain"
    gf = upload_bytes(name, data, mime)
    final = gf
    for _ in range(90):
        try:
            final = client.files.get(name=gf.name)
        except Exception:
            break
        state = getattr(final, "state", None)
        state_text = str(state).upper() if state is not None else ""
        if "FAILED" in state_text:
            raise RuntimeError(f"Gemini failed to process the file: {name}")
        if "PROCESSING" not in state_text:
            break
        time.sleep(1.2)
    return {"name": up.name, "size": up.size, "mime_type": mime, "gemini_file": final}


def render_source(source):
    if source and source.lower() != "not specified in the material":
        st.markdown(f'<span class="source-badge">📌 Source: {source}</span>', unsafe_allow_html=True)


def score_questions(questions, answers):
    correct = 0
    missed = []
    for i, q in enumerate(questions):
        user = str(answers.get(i, "")).strip()
        if user.lower() == q.correct_answer.strip().lower():
            correct += 1
        else:
            missed.append(q)
    return correct, len(questions), missed


def make_questions(count, types_list, difficulties, mode="quiz"):
    if not selected_items():
        raise ValueError("اختاري مادة واحدة على الأقل.")
    task = "Create a realistic university mock exam." if mode == "exam" else "Create an interactive study quiz."
    prompt = f"""
{task}
Number of questions: {count}
Allowed types: {', '.join(types_list)}
Allowed difficulty: {', '.join(difficulties)}

Rules:
- Every question must be supported by the selected uploaded material.
- Cover different important topics and avoid repeating the same fact.
- MCQ must have exactly four options and correct_answer must equal the exact correct option text.
- True/False must have correct_answer exactly True or False and no options.
- Fill in the blank and Essay must have no options.
- Keep answers unambiguous and academically appropriate.
- Add a short explanation grounded in the material.
- Add a short topic label such as "MIS Components", "TPS", "Decision Support", etc.
- Never invent page numbers; use "Not specified in the material" when needed.
"""
    result = ask(prompt, QuestionSet, 0.15)
    qs = result.questions[:count]
    for q in qs:
        if q.question_type == "MCQ" and len(q.options) != 4:
            raise ValueError("تم توليد MCQ غير صالح. جرّبي Generate مرة أخرى.")
    return qs


def make_flashcards(count):
    prompt = f"""
Create {count} concise academic flashcards from the selected material.
Focus on definitions, key terms, important differences, classifications,
processes, formulas, and exam-relevant concepts.
The front should be short. The back must be precise.
Add a short topic label and a reliable source reference when possible.
Never fabricate a page.
"""
    return ask(prompt, FlashcardSet, 0.1).cards[:count]


def render_question(q, i, prefix):
    st.markdown(
        f'<div class="quiz-question"><div class="quiz-number">Question {i+1} · {q.difficulty} · {q.topic}</div><div class="quiz-text">{q.question}</div></div>',
        unsafe_allow_html=True,
    )
    key = f"{prefix}_answer_{i}"
    if q.question_type == "MCQ":
        return st.radio("Choose", q.options, index=None, key=key, label_visibility="collapsed")
    if q.question_type == "True/False":
        return st.radio("Choose", ["True", "False"], index=None, key=key, label_visibility="collapsed")
    return st.text_area("Your answer", key=key, height=100)


def collect_answers(questions, prefix):
    return {i: st.session_state.get(f"{prefix}_answer_{i}", "") for i in range(len(questions))}


def get_weak_topics(questions, answers):
    topics = {}
    for i, q in enumerate(questions):
        user = str(answers.get(i, "")).strip().lower()
        if user != q.correct_answer.strip().lower():
            topics[q.topic] = topics.get(q.topic, 0) + 1
    return sorted(topics.items(), key=lambda x: x[1], reverse=True)

# ============================== SIDEBAR ==============================
with st.sidebar:
    st.markdown('<div class="brand"><div class="brand-icon">🎓</div><div><div class="brand-title">M-MIS Study AI</div><div class="brand-subtitle">Smart Study Assistant</div></div></div>', unsafe_allow_html=True)
    st.caption("Business Administration · Management Information Systems")
    st.divider()
    page = st.radio(
        "Study",
        ["🏠 Dashboard", "💬 AI Tutor", "🧠 Explain & Summarize", "❓ Questions", "🧪 Interactive Quiz", "🗂️ Flashcards", "🎯 Exam Mode"],
        label_visibility="collapsed",
    )
    st.divider()
    st.markdown("### 📚 Materials")
    uploaded = st.file_uploader(
        "Upload study materials",
        type=SUPPORTED_TYPES,
        accept_multiple_files=True,
        help="PDF is best for slides/tables/diagrams. DOCX is extracted to text. You can upload up to 8 files.",
    )
    if uploaded:
        for up in uploaded[:MAX_FILES]:
            if up.name not in st.session_state.materials:
                with st.spinner(f"Uploading {up.name}..."):
                    try:
                        st.session_state.materials[up.name] = process_upload(up)
                        save_action("Upload material", len(st.session_state.materials))
                    except Exception as exc:
                        st.error(f"تعذر رفع {up.name}: {exc}")
        if not st.session_state.selected_materials:
            st.session_state.selected_materials = list(st.session_state.materials.keys())

    if st.session_state.materials:
        names = list(st.session_state.materials.keys())
        selected = st.multiselect(
            "Select materials for AI",
            names,
            default=[n for n in st.session_state.selected_materials if n in names],
        )
        if selected != st.session_state.selected_materials:
            st.session_state.selected_materials = selected
            reset_study_state()
        for n in selected:
            size = st.session_state.materials[n]["size"] / 1048576
            st.markdown(f'<div class="material-chip"><span>📄</span><b title="{n}">{n}</b><span style="margin-left:auto;color:#778196">{size:.1f} MB</span></div>', unsafe_allow_html=True)
        if st.button("🗑️ Clear materials", use_container_width=True):
            st.session_state.materials = {}
            st.session_state.selected_materials = []
            reset_study_state()
            st.rerun()
    else:
        st.info("ارفعي أول محاضرة للبدء.")
    st.divider()
    st.markdown('<div class="status"><span class="status-dot"></span>Gemini 3.8 Flash · Interactions API</div>', unsafe_allow_html=True)

# ============================== HERO ==============================
count_selected = len(selected_items())
st.markdown(
    f'<div class="hero"><div class="hero-kicker">🎓 M-MIS Study AI</div><h1>Study smarter. Know your material.</h1><p>ارفعي محاضراتك، اسألي عنها، لخصيها، افهميها، اعملي Quiz وFlashcards وMock Exam — مع تركيز على الدقة والرجوع للمادة الأصلية.</p><div class="status"><span class="status-dot"></span>{count_selected} selected material(s)</div></div>',
    unsafe_allow_html=True,
)

# ============================== DASHBOARD ==============================
if page == "🏠 Dashboard":
    history = get_history(8)
    total_attempts = len(history)
    avg = round(sum(r[3] for r in history) / total_attempts) if total_attempts else 0
    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown(f'<div class="metric"><div class="metric-number">{count_selected}</div><div class="metric-label">Materials selected</div></div>', unsafe_allow_html=True)
    with c2:
        st.markdown(f'<div class="metric"><div class="metric-number">{total_attempts}</div><div class="metric-label">Saved quiz/exam attempts</div></div>', unsafe_allow_html=True)
    with c3:
        st.markdown(f'<div class="metric"><div class="metric-number">{avg}%</div><div class="metric-label">Recent average</div></div>', unsafe_allow_html=True)

    st.markdown('<div class="section-title">✨ Study Tools</div>', unsafe_allow_html=True)
    features = [
        ("💬", "AI Tutor", "اسألي عن المحاضرات المختارة مع محادثة مستمرة."),
        ("🧠", "Explain", "شرح بالعربي مع المصطلحات الإنجليزية المهمة."),
        ("📝", "Summarize", "ملخص منظم للنقاط المهمة والتعريفات."),
        ("❓", "Questions", "MCQ وTrue/False وFill in وEssay."),
        ("🧪", "Interactive Quiz", "جاوبي واحسبي الدرجة تلقائيًا."),
        ("🗂️", "Flashcards", "مراجعة سريعة للمفاهيم الأساسية."),
        ("🎯", "Exam Mode", "Mock Exam + تحليل نقاط الضعف."),
    ]
    cols = st.columns(3)
    for i, (icon, title, text) in enumerate(features):
        with cols[i % 3]:
            st.markdown(f'<div class="card feature-card"><div class="feature-icon">{icon}</div><div class="feature-title">{title}</div><div class="feature-text">{text}</div></div><br>', unsafe_allow_html=True)

    st.markdown('<div class="section-title">📈 Recent Performance</div>', unsafe_allow_html=True)
    if history:
        import pandas as pd
        df = pd.DataFrame(history, columns=["Mode", "Score", "Total", "Percent", "Date"])
        st.dataframe(df, use_container_width=True, hide_index=True)
    else:
        st.info("اعملي أول Quiz أو Mock Exam وهيظهر هنا.")

    st.warning("Accuracy first: الذكاء الاصطناعي ممكن يخطئ. التعريفات والنقط المهمة جدًا راجعيها مع الـofficial course material قبل الامتحان.")

# ============================== AI TUTOR ==============================
elif page == "💬 AI Tutor":
    st.markdown('<div class="section-title">💬 AI Tutor</div>', unsafe_allow_html=True)
    if not selected_items():
        st.info("📚 اختاري مادة واحدة على الأقل.")
    else:
        for m in st.session_state.chat_messages:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])
        prompt = st.chat_input("مثال: Explain the difference between TPS and MIS.")
        if prompt:
            st.session_state.chat_messages.append({"role": "user", "content": prompt})
            with st.chat_message("user"):
                st.markdown(prompt)
            tutor_prompt = f"""
The student asks: {prompt}

Answer directly using the selected uploaded materials as the primary source.
Give the answer first, then a short explanation and key English terms.
If the answer is not clearly supported by the material, say exactly:
This information is not clearly stated in the uploaded material.
If general academic knowledge is needed, label it: General academic clarification.
Never fabricate a page number.
"""
            with st.chat_message("assistant"):
                with st.spinner("🧠 Analyzing..."):
                    try:
                        ans = ask(tutor_prompt, temperature=0.1, continue_chat=True)
                        st.markdown(ans)
                        st.session_state.chat_messages.append({"role": "assistant", "content": ans})
                    except Exception as exc:
                        st.error(f"حدث خطأ: {exc}")

# ============================== EXPLAIN ==============================
elif page == "🧠 Explain & Summarize":
    st.markdown('<div class="section-title">🧠 Learn the Material</div>', unsafe_allow_html=True)
    if not selected_items():
        st.info("📚 اختاري مادة أولًا.")
    else:
        action = st.radio("Choose", ["Explain", "Summarize", "Key Points", "Exam Cram"], horizontal=True)
        if st.button(f"✨ Generate {action}", type="primary", use_container_width=True):
            prompts = {
                "Explain": "Explain the selected material in Arabic step by step. Cover the big picture, concepts, definitions, processes, differences, supported examples, and exam takeaways.",
                "Summarize": "Create an accurate academic summary of the selected material covering main topics, definitions, concepts, classifications, processes, relationships, examples, and exam-focused points.",
                "Key Points": "Extract concise exam-relevant points. Prioritize definitions, English terms, classifications, differences, processes, formulas, and relationships.",
                "Exam Cram": "Create a last-minute exam revision sheet: must-memorize definitions, key terms, comparisons, classifications, processes, formulas, common confusions, and likely exam points. Only use supported material.",
            }
            with st.spinner("جاري تجهيز المحتوى..."):
                try:
                    st.markdown(ask(prompts[action], temperature=0.1))
                    save_action(action, len(selected_items()))
                except Exception as exc:
                    st.error(f"حدث خطأ: {exc}")

# ============================== QUESTIONS ==============================
elif page == "❓ Questions":
    st.markdown('<div class="section-title">❓ Generate Questions</div>', unsafe_allow_html=True)
    if not selected_items():
        st.info("📚 اختاري مادة أولًا.")
    else:
        c1, c2 = st.columns(2)
        with c1:
            qcount = st.slider("Number of questions", 5, 30, 10)
            qtypes = st.multiselect("Question types", ["MCQ", "True/False", "Fill in the blank", "Essay"], default=["MCQ", "True/False"])
        with c2:
            diffs = st.multiselect("Difficulty", ["Easy", "Medium", "Hard"], default=["Easy", "Medium", "Hard"])
        if st.button("🚀 Generate Question Set", type="primary", use_container_width=True):
            if not qtypes or not diffs:
                st.warning("اختاري نوع سؤال ومستوى صعوبة.")
            else:
                with st.spinner("🤖 Generating questions..."):
                    try:
                        st.session_state.quiz_questions = make_questions(qcount, qtypes, diffs)
                        save_action("Generate questions", len(selected_items()))
                        st.success(f"تم إنشاء {len(st.session_state.quiz_questions)} سؤال.")
                    except Exception as exc:
                        st.error(f"تعذر إنشاء الأسئلة: {exc}")
        if st.session_state.quiz_questions:
            for i, q in enumerate(st.session_state.quiz_questions):
                st.markdown(f"**{i+1}. {q.question}**  \n`{q.question_type}` · `{q.difficulty}` · `{q.topic}`")
                if q.question_type == "MCQ":
                    st.markdown("\n".join(f"- **{chr(65+j)}.** {o}" for j, o in enumerate(q.options)))
                st.markdown(f"**Answer:** {q.correct_answer}")
                st.caption(q.explanation)
                render_source(q.source)
                st.divider()

# ============================== QUIZ ==============================
elif page == "🧪 Interactive Quiz":
    st.markdown('<div class="section-title">🧪 Interactive Quiz</div>', unsafe_allow_html=True)
    if not selected_items():
        st.info("📚 اختاري مادة أولًا.")
    else:
        if not st.session_state.quiz_questions:
            if st.button("⚡ Create 10-question Quiz", type="primary", use_container_width=True):
                with st.spinner("جاري إنشاء الـQuiz..."):
                    try:
                        st.session_state.quiz_questions = make_questions(10, ["MCQ", "True/False"], ["Easy", "Medium", "Hard"])
                        st.session_state.quiz_answers = {}
                        st.session_state.quiz_submitted = False
                        st.session_state.quiz_score = None
                        save_action("Create quiz", len(selected_items()))
                        st.rerun()
                    except Exception as exc:
                        st.error(f"حدث خطأ: {exc}")
        else:
            if not st.session_state.quiz_submitted:
                for i, q in enumerate(st.session_state.quiz_questions):
                    render_question(q, i, "quiz")
                if st.button("✅ Submit Quiz", type="primary", use_container_width=True):
                    st.session_state.quiz_answers = collect_answers(st.session_state.quiz_questions, "quiz")
                    correct, total, _ = score_questions(st.session_state.quiz_questions, st.session_state.quiz_answers)
                    st.session_state.quiz_score = correct
                    st.session_state.quiz_submitted = True
                    topics = [t for t, _ in get_weak_topics(st.session_state.quiz_questions, st.session_state.quiz_answers)]
                    save_result("Quiz", correct, total, topics)
                    st.rerun()
            else:
                correct = st.session_state.quiz_score or 0
                total = len(st.session_state.quiz_questions)
                pct = round(correct / total * 100) if total else 0
                st.markdown(f'<div class="score-card"><div style="color:#a5b4fc;font-weight:700">YOUR SCORE</div><div class="score-number">{pct}%</div><div style="color:#aeb7ca">{correct} / {total} correct</div></div>', unsafe_allow_html=True)
                weak = get_weak_topics(st.session_state.quiz_questions, st.session_state.quiz_answers)
                if weak:
                    st.markdown("### 🎯 Weak Topics")
                    st.write(" · ".join(f"{t} ({n})" for t, n in weak[:5]))
                for i, q in enumerate(st.session_state.quiz_questions):
                    user = st.session_state.quiz_answers.get(i, "")
                    if str(user).strip().lower() == q.correct_answer.strip().lower():
                        st.success(f"Q{i+1}: Correct — {q.correct_answer}")
                    else:
                        st.error(f"Q{i+1}: Your answer: {user or 'No answer'}")
                        st.markdown(f"**Correct answer:** {q.correct_answer}")
                        st.caption(q.explanation)
                        render_source(q.source)
                if st.button("🔄 New Quiz", use_container_width=True):
                    st.session_state.quiz_questions = []
                    st.session_state.quiz_answers = {}
                    st.session_state.quiz_submitted = False
                    st.session_state.quiz_score = None
                    st.rerun()

# ============================== FLASHCARDS ==============================
elif page == "🗂️ Flashcards":
    st.markdown('<div class="section-title">🗂️ Flashcards</div>', unsafe_allow_html=True)
    if not selected_items():
        st.info("📚 اختاري مادة أولًا.")
    else:
        if not st.session_state.flashcards:
            n = st.slider("Number of flashcards", 5, 30, 10)
            if st.button("✨ Generate Flashcards", type="primary", use_container_width=True):
                with st.spinner("جاري إنشاء Flashcards..."):
                    try:
                        st.session_state.flashcards = make_flashcards(n)
                        st.session_state.flashcard_index = 0
                        st.session_state.flashcard_revealed = False
                        save_action("Generate flashcards", len(selected_items()))
                        st.rerun()
                    except Exception as exc:
                        st.error(f"حدث خطأ: {exc}")
        if st.session_state.flashcards:
            cards = st.session_state.flashcards
            i = st.session_state.flashcard_index
            card = cards[i]
            st.progress((i + 1) / len(cards), text=f"Card {i+1} of {len(cards)} · {card.topic}")
            if not st.session_state.flashcard_revealed:
                st.markdown(f'<div class="card" style="min-height:260px;text-align:center;display:flex;flex-direction:column;justify-content:center"><div style="color:#8d96aa">QUESTION</div><div style="font-size:30px;font-weight:850">{card.front}</div></div>', unsafe_allow_html=True)
                if st.button("👀 Reveal Answer", type="primary", use_container_width=True):
                    st.session_state.flashcard_revealed = True
                    st.rerun()
            else:
                st.markdown(f'<div class="card" style="min-height:260px;text-align:center;display:flex;flex-direction:column;justify-content:center"><div style="color:#8d96aa">ANSWER</div><div style="font-size:24px;font-weight:750;line-height:1.5">{card.back}</div></div>', unsafe_allow_html=True)
                render_source(card.source)
                a, b, c = st.columns(3)
                with a:
                    if st.button("← Previous", use_container_width=True):
                        st.session_state.flashcard_index = max(0, i - 1)
                        st.session_state.flashcard_revealed = False
                        st.rerun()
                with b:
                    if st.button("🔁 Again", use_container_width=True):
                        st.session_state.flashcard_revealed = False
                        st.rerun()
                with c:
                    if i < len(cards) - 1:
                        if st.button("Next →", use_container_width=True):
                            st.session_state.flashcard_index += 1
                            st.session_state.flashcard_revealed = False
                            st.rerun()
                    else:
                        st.success("🎉 Deck complete!")
            if st.button("🗑️ Generate a new deck", use_container_width=True):
                st.session_state.flashcards = []
                st.session_state.flashcard_index = 0
                st.session_state.flashcard_revealed = False
                st.rerun()

# ============================== EXAM ==============================
elif page == "🎯 Exam Mode":
    st.markdown('<div class="section-title">🎯 Exam Mode</div>', unsafe_allow_html=True)
    if not selected_items():
        st.info("📚 اختاري مادة أولًا.")
    else:
        st.markdown('<div class="card"><div class="feature-title">🎓 Mock Exam</div><div class="feature-text">امتحان تجريبي مبني على المادة المختارة، مع تصحيح تلقائي وتحليل لنقاط الضعف.</div></div>', unsafe_allow_html=True)
        if not st.session_state.exam_questions:
            c1, c2 = st.columns(2)
            with c1:
                n = st.slider("Exam questions", 5, 30, 15)
            with c2:
                diffs = st.multiselect("Difficulty", ["Easy", "Medium", "Hard"], default=["Medium", "Hard"])
            if st.button("🚀 Start Mock Exam", type="primary", use_container_width=True):
                if not diffs:
                    st.warning("اختاري مستوى صعوبة.")
                else:
                    with st.spinner("جاري إعداد الـMock Exam..."):
                        try:
                            st.session_state.exam_questions = make_questions(n, ["MCQ", "True/False", "Fill in the blank"], diffs, "exam")
                            st.session_state.exam_answers = {}
                            st.session_state.exam_submitted = False
                            st.session_state.exam_score = None
                            st.session_state.exam_started_at = time.time()
                            save_action("Start mock exam", len(selected_items()))
                            st.rerun()
                        except Exception as exc:
                            st.error(f"تعذر إنشاء الامتحان: {exc}")
        else:
            if not st.session_state.exam_submitted:
                if st.session_state.exam_started_at:
                    elapsed = int(time.time() - st.session_state.exam_started_at)
                    mins, secs = divmod(elapsed, 60)
                    st.info(f"⏱️ Time elapsed: {mins:02d}:{secs:02d}")
                for i, q in enumerate(st.session_state.exam_questions):
                    render_question(q, i, "exam")
                if st.button("📝 Submit Exam", type="primary", use_container_width=True):
                    st.session_state.exam_answers = collect_answers(st.session_state.exam_questions, "exam")
                    correct, total, _ = score_questions(st.session_state.exam_questions, st.session_state.exam_answers)
                    st.session_state.exam_score = correct
                    st.session_state.exam_submitted = True
                    weak_topics = [t for t, _ in get_weak_topics(st.session_state.exam_questions, st.session_state.exam_answers)]
                    save_result("Mock Exam", correct, total, weak_topics)
                    st.rerun()
            else:
                correct = st.session_state.exam_score or 0
                total = len(st.session_state.exam_questions)
                pct = round(correct / total * 100) if total else 0
                st.markdown(f'<div class="score-card"><div style="color:#a5b4fc;font-weight:700">MOCK EXAM SCORE</div><div class="score-number">{pct}%</div><div style="color:#aeb7ca">{correct} / {total} correct</div></div>', unsafe_allow_html=True)
                weak = get_weak_topics(st.session_state.exam_questions, st.session_state.exam_answers)
                if weak:
                    st.markdown("### 🎯 Topics to Review")
                    for topic, count in weak[:6]:
                        st.warning(f"{topic}: {count} missed")
                else:
                    st.success("🎉 No incorrect answers. Great job!")
                st.markdown("### 📚 Detailed Review")
                for i, q in enumerate(st.session_state.exam_questions):
                    with st.expander(f"Question {i+1}: {q.question}"):
                        st.write(f"Your answer: {st.session_state.exam_answers.get(i,'') or 'No answer'}")
                        st.write(f"Correct answer: {q.correct_answer}")
                        st.write(q.explanation)
                        render_source(q.source)
                if st.button("🔄 Create New Mock Exam", use_container_width=True):
                    st.session_state.exam_questions = []
                    st.session_state.exam_answers = {}
                    st.session_state.exam_submitted = False
                    st.session_state.exam_score = None
                    st.session_state.exam_started_at = None
                    st.rerun()

st.divider()
st.caption("M-MIS Study AI · Built for study support · Gemini 3.8 Flash · Verify critical exam information against your official course material.")
